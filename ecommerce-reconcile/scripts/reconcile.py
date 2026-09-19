"""Fixed-template reconciliation. Integer cents are the only accounting unit."""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import openpyxl

VERSION = '1.0.0'
SPECS = {
    'orders': ('订单.xlsx', ['订单号', '付款日期', '实收金额', '币种']),
    'refunds': ('退款.xlsx', ['退款单号', '订单号', '退款日期', '退款金额', '退款状态', '币种']),
    'fees': ('平台费用.xlsx', ['费用单号', '订单号', '发生日期', '费用金额', '费用类型', '币种']),
    'costs': ('成本.xlsx', ['订单号', '商品成本', '履约成本', '成本冲回', '币种']),
}
MONEY = {'实收金额', '退款金额', '费用金额', '商品成本', '履约成本', '成本冲回'}
DATES = {'付款日期', '退款日期', '发生日期'}
MAX_CENTS = 10_000_000_000  # 单个金额最多 1 亿元；最多 1 万行，汇总保持 Excel 精度。
MAX_ROWS = 10_000
TIPS = {
    'MISSING_FILE': '补交缺失文件，保留文件名和模板表头。',
    'FILE_INVALID': '提供未加密、可打开的标准 xlsx 文件。',
    'SCHEMA': '使用四表模板，保留唯一工作表“数据”和全部规定表头。',
    'EMPTY_ORDERS': '至少提供一笔已成功收款的订单。',
    'INVALID_CELL': '核实原始凭证，在新副本中修正后重新运行。禁止直接改报表。',
    'DUPLICATE_ID': '核实重复记录的来源；不能自动保留第一条或重复累计。',
    'ORPHAN_ORDER': '核实订单号，补交关联订单或单独划定批次。',
    'MISSING_COST': '补交该订单成本；确实为零也须明确填 0。',
    'MISSING_FEE': '补交该订单费用；确认无费用时提供金额为 0 的记录。',
    'REFUND_EXCEEDS_PAID': '核实实收及全部成功退款，排查重复退款或口径不一致。',
    'RECOVERY_EXCEEDS_COST': '成本冲回不能大于已列商品成本。',
    'DATE_BEFORE_PAID': '核实付款日期和退款或费用发生日期。',
    'LARGE_AMOUNT': '金额达到预设提醒阈值，请对照凭证复核。',
    'FEE_EXCEEDS_PAID': '费用大于订单实收，请复核费用归属及金额。',
    'NEGATIVE_PROFIT': '该订单经营贡献为负，请核查成本、费用与退款。',
    'REFUND_PENDING': '尚未成功的退款未扣减，后续成功时请更新副本重跑。',
    'REFUND_FAILED': '失败退款未扣减，核实后续是否重新发起。',
}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def money(value):
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError('金额须为数值或普通十进制文本')
    s = str(value)
    if not re.fullmatch(r'\d+(?:\.\d{1,2})?', s):
        raise ValueError('金额必须非负且最多两位小数，不接受空格、千分位、科学计数文本或货币符号')
    try:
        d = Decimal(s)
    except InvalidOperation as exc:
        raise ValueError('金额格式无效') from exc
    if not d.is_finite() or d > Decimal(MAX_CENTS) / 100:
        raise ValueError('单个金额超过 1 亿元上限')
    return int(d * 100)


def identifier(value):
    if not isinstance(value, str) or not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,63}', value):
        raise ValueError('编号必须为 1–64 位英文、数字、下划线或连字符文本；数字单元格不接受')
    return value


def calendar_date(value):
    if isinstance(value, datetime):
        if value.time().isoformat() != '00:00:00':
            raise ValueError('日期不能包含非零时分秒')
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        raise ValueError('日期须为 YYYY-MM-DD 或 Excel 日期单元格')
    return date.fromisoformat(value).isoformat()


def raw_text(value):
    if value is None:
        return ''
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value)


def calculate(input_dir, as_of, high_amount_cents=10_000_000, archive=None):
    """Read-only calculation. Returns source-row evidence, exclusions and exact cents."""
    as_of = calendar_date(as_of)
    sources, ledger, issues = {}, {}, []
    fatal = False

    def issue(code, kind, row=0, field='', value='', order='', detail='', severity='阻断'):
        issues.append(dict(code=code, severity=severity, file=SPECS[kind][0], sheet='数据',
                           row=row, field=field, raw_value=raw_text(value), order_id=order,
                           detail=detail, action=TIPS[code]))

    for kind, (filename, headers) in SPECS.items():
        path = Path(input_dir) / filename
        ledger[kind] = []
        if not path.is_file():
            issue('MISSING_FILE', kind, detail='文件缺失，不能判断相关金额是否完整')
            fatal = True
            continue
        content = path.read_bytes()
        digest = sha(content)
        sources[kind] = dict(file=filename, source_path=str(path.resolve()), sha256=digest,
                             bytes=len(content), blank_rows=[], data_rows=0)
        if archive:
            Path(archive, filename).write_bytes(content)
        book = None
        try:
            if len(content) > 20 * 1024 * 1024:
                raise ValueError('文件超过 20 MB')
            import io
            with ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('解压后内容超过 100 MB')
            book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            if book.sheetnames != ['数据']:
                issue('SCHEMA', kind, detail='仅允许一个名为“数据”的工作表')
                fatal = True
                continue
            ws = book['数据']
            ws.reset_dimensions()  # Read actual cells even when an XLSX omits or understates dimensions.
            ws.calculate_dimension(force=True)
            if ws.max_row > MAX_ROWS + 1 or ws.max_column > 32:
                raise ValueError('超出 MVP 上限：每表 10000 行，最多 32 列（含已格式化区域）')
            rows = list(ws.iter_rows())
            actual = [c.value for c in rows[0]] if rows else []
            while actual and actual[-1] is None:
                actual.pop()
            schema_ok = len(actual) == len(headers) and set(actual) == set(headers)
            if not schema_ok:
                issue('SCHEMA', kind, row=1, value=json.dumps(actual, ensure_ascii=False, default=str),
                      detail='表头缺失、重复、多余或拼写不一致；该表不参与计算')
                fatal = True
            for rn, cells in enumerate(rows[1:], 2):
                if all(c.value is None for c in cells):
                    sources[kind]['blank_rows'].append(rn)
                    continue
                r = dict(row=rn, raw={c.coordinate: raw_text(c.value) for c in cells if c.value is not None},
                         values={}, valid=schema_ok, disposition='待处理', order_id='')
                ledger[kind].append(r)
                if not schema_ok:
                    r['disposition'] = '表结构无效'
                    continue
                for field, cell in zip(actual, cells):
                    v = cell.value
                    try:
                        if cell.data_type in ('f', 'e'):
                            raise ValueError('不接受 Excel 公式或错误值；请提供已核实的值副本')
                        if v is None or v == '':
                            raise ValueError('必填字段为空')
                        if field.endswith('号'):
                            val = identifier(v)
                        elif field in MONEY:
                            val = money(v)
                            if field == '实收金额' and val == 0:
                                raise ValueError('本模板仅接受实收大于 0 的成功付款订单')
                        elif field in DATES:
                            val = calendar_date(v)
                            if val > as_of:
                                raise ValueError('日期晚于核对截止日')
                        elif field == '币种':
                            if v != 'CNY':
                                raise ValueError('本模板仅支持 CNY，不做汇率换算')
                            val = v
                        elif field == '退款状态':
                            if v not in ('成功', '处理中', '失败'):
                                raise ValueError('退款状态只接受：成功、处理中、失败')
                            val = v
                        else:
                            if v not in ('平台佣金', '支付手续费', '其他订单费用', '无费用'):
                                raise ValueError('费用类型不在模板允许范围')
                            val = v
                        r['values'][field] = val
                    except (ValueError, TypeError) as exc:
                        r['valid'] = False
                        issue('INVALID_CELL', kind, rn, field, v, detail=str(exc))
                r['order_id'] = r['values'].get('订单号', '')
                if kind == 'fees' and r['values'].get('费用类型') == '无费用' and r['values'].get('费用金额', 0) != 0:
                    r['valid'] = False
                    issue('INVALID_CELL', kind, rn, '费用金额', r['values']['费用金额'], r['order_id'], '“无费用”记录金额必须为 0（此处原值单位为分）')
                # Extra data beyond headers is never silently discarded.
                if any(c.value is not None for c in cells[len(actual):]):
                    r['valid'] = False
                    issue('SCHEMA', kind, rn, order=r['order_id'], detail='规定列以外存在数据')
                    fatal = True
            sources[kind]['data_rows'] = len(ledger[kind])
        except Exception as exc:
            # Corrupt third-party XLSX readers raise several distinct exception classes.
            issue('FILE_INVALID', kind, detail=f'{type(exc).__name__}: {exc}')
            fatal = True
        finally:
            if book:
                book.close()

    # Attach order IDs to validation findings only after the full row has been parsed.
    for kind, rows in ledger.items():
        by_row = {r['row']: r for r in rows}
        for i in issues:
            if i['file'] == SPECS[kind][0] and i['row'] in by_row:
                i['order_id'] = by_row[i['row']]['order_id']
        key = {'orders': '订单号', 'costs': '订单号', 'refunds': '退款单号', 'fees': '费用单号'}[kind]
        groups = defaultdict(list)
        for r in rows:
            if key in r['values']:
                groups[r['values'][key]].append(r)
        for value, matches in groups.items():
            if len(matches) > 1:
                for r in matches:
                    r['valid'] = False
                    issue('DUPLICATE_ID', kind, r['row'], key, value, r['order_id'], '同一编号的全部记录均隔离，不自动去重')

    order_groups = defaultdict(list)
    for r in ledger['orders']:
        if r['order_id']:
            order_groups[r['order_id']].append(r)
    if not ledger['orders']:
        issue('EMPTY_ORDERS', 'orders', detail='没有可识别的非空订单记录')
        fatal = True
    related = {k: defaultdict(list) for k in ('refunds', 'fees', 'costs')}
    for kind in related:
        for r in ledger[kind]:
            oid = r['order_id']
            related[kind][oid].append(r)
            if oid and oid not in order_groups:
                r['valid'] = False
                issue('ORPHAN_ORDER', kind, r['row'], '订单号', oid, oid, '无法在订单表找到关联订单')

    details = []
    for oid, group in order_groups.items():
        order = group[0]
        v = order['values']
        children = {k: related[k][oid] for k in related}
        for kind, code in [('costs', 'MISSING_COST'), ('fees', 'MISSING_FEE')]:
            if not children[kind]:
                issue(code, kind, order=oid, detail='缺少明确记录，不能按 0 计算')
        all_valid = len(group) == 1 and all(r['valid'] for r in group) and all(
            r['valid'] for rows in children.values() for r in rows)
        if all_valid and children['costs'] and children['fees'] and not fatal:
            for kind, field in [('refunds', '退款日期'), ('fees', '发生日期')]:
                for r in children[kind]:
                    if r['values'][field] < v['付款日期']:
                        issue('DATE_BEFORE_PAID', kind, r['row'], field, r['values'][field], oid, '发生日期早于付款日期')
            refund = sum(r['values']['退款金额'] for r in children['refunds'] if r['values']['退款状态'] == '成功')
            fee = sum(r['values']['费用金额'] for r in children['fees'])
            cost = children['costs'][0]['values']
            if refund > v['实收金额']:
                issue('REFUND_EXCEEDS_PAID', 'refunds', order=oid, detail=f'成功退款 {refund} 分 > 实收 {v["实收金额"]} 分')
            if cost['成本冲回'] > cost['商品成本']:
                issue('RECOVERY_EXCEEDS_COST', 'costs', children['costs'][0]['row'], order=oid, detail='成本冲回超过商品成本')
            blockers = [i['code'] for i in issues if i['order_id'] == oid and i['severity'] == '阻断']
            if not blockers:
                amounts = dict(paid=v['实收金额'], refund=refund, fee=fee, goods=cost['商品成本'],
                               fulfillment=cost['履约成本'], recovery=cost['成本冲回'])
                amounts['profit'] = amounts['paid'] - refund - fee - amounts['goods'] - amounts['fulfillment'] + amounts['recovery']
                if amounts['paid'] >= high_amount_cents:
                    issue('LARGE_AMOUNT', 'orders', order['row'], order=oid, detail=f'实收达到提醒阈值 {high_amount_cents} 分', severity='提醒')
                if fee > v['实收金额']:
                    issue('FEE_EXCEEDS_PAID', 'fees', order=oid, detail='费用高于实收，金额仍按原值计入', severity='提醒')
                if amounts['profit'] < 0:
                    issue('NEGATIVE_PROFIT', 'orders', order['row'], order=oid, detail='订单经营贡献为负，金额仍按原值计入', severity='提醒')
            else:
                amounts = None
        else:
            amounts = None
        for r in children['refunds']:
            if r['valid'] and r['values'].get('退款状态') in ('处理中', '失败'):
                status = r['values']['退款状态']
                issue('REFUND_PENDING' if status == '处理中' else 'REFUND_FAILED', 'refunds', r['row'],
                      '退款状态', status, oid, '未作为成功退款扣减', '提醒' if status == '处理中' else '信息')
        blockers = sorted({i['code'] for i in issues if i['order_id'] == oid and i['severity'] == '阻断'})
        if fatal:
            blockers.append('全批次文件或表结构不完整')
        details.append(dict(order_id=oid, source_rows=[r['row'] for r in group],
                            payment_date=v.get('付款日期'), status='可计算' if amounts is not None else '待核实',
                            amounts_cents=amounts, blockers=blockers,
                            references={k: [r['row'] for r in rows] for k, rows in children.items()}))

    eligible = {d['order_id'] for d in details if d['amounts_cents'] is not None}
    for kind, rows in ledger.items():
        for r in rows:
            if not r['valid'] or r['order_id'] not in eligible:
                r['disposition'] = '隔离待核实'
            elif kind == 'refunds' and r['values']['退款状态'] != '成功':
                r['disposition'] = '未成功退款，不计入'
            else:
                r['disposition'] = '计入'
    totals = {k: sum(d['amounts_cents'][k] for d in details if d['amounts_cents'] is not None)
              for k in ('paid', 'refund', 'fee', 'goods', 'fulfillment', 'recovery', 'profit')}
    counts = Counter(i['severity'] for i in issues)
    complete = not counts['阻断'] and bool(details)
    dates = [d['payment_date'] for d in details if d['payment_date']]
    return dict(version=VERSION, as_of=as_of, currency='CNY', high_amount_cents=high_amount_cents,
                status='完整（有提醒）' if complete and counts['提醒'] else '完整' if complete else '不完整，禁止作为全店利润',
                complete=complete, sources=sources, ledger=ledger, issues=issues, details=details,
                totals_cents=totals, eligible_orders=len(eligible), unique_orders=len(order_groups),
                source_order_rows=len(ledger['orders']), missing_id_order_rows=sum(not r['order_id'] for r in ledger['orders']),
                issue_counts=dict(counts), payment_range=[min(dates), max(dates)] if dates else [],
                row_counts={k: dict(Counter(r['disposition'] for r in rows)) for k, rows in ledger.items()})


def rules(result):
    return [
        ['结果状态', result['status']],
        ['口径', '订单经营贡献；按本批次订单归集截至核对日的已提供流水。不是自然月权责利润或银行到账核对。'],
        ['计算式', '实收 − 成功退款 − 平台费用 − 商品成本 − 履约成本 + 成本冲回'],
        ['单位与精度', 'CNY 元，输入最多两位小数。Decimal 解析为整数分后加减，不舍入，不使用 AI 生成金额。'],
        ['实收定义', '原始成功收款，含买家支付运费、已扣优惠、未扣退款与平台费用。不得填退款后净收款。'],
        ['退款', '仅“成功”扣减。每笔退款必须有唯一编号；处理中或失败单独列示。退款无记录按提供范围内无退款处理。'],
        ['费用', '每个订单至少一行费用。无费用也要用唯一费用单号明确填 0；费用为已扣或已确认的订单费用。'],
        ['成本', '每个订单恰好一行。商品成本为本订单总成本，不是单价；履约成本包含商家承担的运费包装等。'],
        ['成本冲回', '由商家凭证明确提供，0 也必须填写；退货或退款不会自动触发成本冲回。上限为商品成本。'],
        ['隔离规则', '重复编号的全部记录、缺失或无效金额及其关联订单不计入小计。文件或表结构异常阻断全批次计算。'],
        ['缺失关联', '无法关联的退款、费用或成本会使整体不完整；可计算订单小计仅供逐单核对，不能代替全店利润。'],
        ['异常金额', f'实收达到 {result["high_amount_cents"] // 100} 元、费用大于实收或经营贡献为负会提醒；原值仍计入。退款超过实收则阻断。'],
        ['完整性的范围', '“完整”仅表示已上传四表通过规则校验，不能证明商家没有漏上传交易。商家须确认四表覆盖同一订单批次。'],
        ['期间', f'核对截止日 {result["as_of"]}；不按自然月截断关联退款和费用。晚于截止日或早于付款日的流水需核实。'],
        ['保留原始数据', '原始文件副本逐字节保存并记录 SHA-256；全部非空行均在审计 JSON 中记录去向，空白行单独登记。'],
        ['查看与重跑', 'Excel 为确定性计算快照，改动报表不会重新计算。修正来源的新副本后，以新的输出目录重跑。'],
        ['未覆盖项目', '多币种、负数费用冲销、SKU 拆单、广告租金人工税费分摊、库存计价、发票和结算到账均不在本版范围。'],
        ['容量', '四个 .xlsx，每表最多 10000 行、每个文件最多 20 MB；只支持固定四表 v1。'],
    ]


def find_node():
    found = shutil.which('node')
    bundled = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
    return str(bundled) if bundled.is_file() else found or 'node'


def run(input_dir, output_dir, as_of, high_amount_cents=10_000_000, render=False):
    input_dir, output_dir = Path(input_dir).resolve(), Path(output_dir).resolve()
    if not input_dir.is_dir():
        raise ValueError('输入目录不存在')
    if output_dir == input_dir or output_dir.is_relative_to(input_dir):
        raise ValueError('输出目录必须与输入目录分开')
    output_dir.mkdir(parents=True, exist_ok=False)  # Never overwrite an earlier audit run.
    raw = output_dir / '原始数据'
    raw.mkdir()
    result = calculate(input_dir, as_of, high_amount_cents, raw)
    result['created_utc'] = datetime.now(timezone.utc).isoformat()
    result['code_sha256'] = {p.name: sha(p.read_bytes()) for p in (Path(__file__), Path(__file__).with_name('export.mjs'))}
    result['rules'] = rules(result)
    for source in result['sources'].values():
        if sha(Path(source['source_path']).read_bytes()) != source['sha256']:
            raise RuntimeError('运行期间输入文件被修改，请保留本次目录并从稳定副本重跑')
    (output_dir / '审计记录.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    (output_dir / '异常清单.json').write_text(json.dumps(result['issues'], ensure_ascii=False, indent=2), encoding='utf-8')
    explanation = '# 计算说明\n\n' + '\n\n'.join(f'**{k}**：{v}' for k, v in result['rules'])
    (output_dir / '计算说明.md').write_text(explanation + '\n', encoding='utf-8')
    cmd = [find_node(), str(Path(__file__).with_name('export.mjs')), 'report', str(output_dir / '审计记录.json')]
    if render:
        cmd.append('--render')
    subprocess.run(cmd, check=True)
    artifacts = {p.name: sha(p.read_bytes()) for p in output_dir.iterdir() if p.is_file()}
    (output_dir / '交付校验.json').write_text(json.dumps(artifacts, ensure_ascii=False, indent=2), encoding='utf-8')
    (output_dir / '完成标记.txt').write_text('报表及审计文件已导出。业务状态：' + result['status'], encoding='utf-8')
    return result


def main():
    parser = argparse.ArgumentParser(description='固定四表 v1 电商经营核对；输出为独立审计快照')
    parser.add_argument('--input', required=True, help='四个固定文件所在目录')
    parser.add_argument('--output', required=True, help='必须为尚不存在的新目录')
    parser.add_argument('--as-of', required=True, help='核对截止日 YYYY-MM-DD')
    parser.add_argument('--high-amount', default='100000.00', help='实收大额提醒阈值，元；默认 100000.00')
    parser.add_argument('--render', action='store_true', help='额外生成工作表预览，用于演示验收')
    args = parser.parse_args()
    try:
        threshold = money(args.high_amount)
        if threshold <= 0:
            raise ValueError('提醒阈值必须大于 0')
        result = run(args.input, args.output, calendar_date(args.as_of), threshold, args.render)
        print(json.dumps({'status': result['status'], 'eligible_orders': result['eligible_orders'],
                          'profit_cents': result['totals_cents']['profit'], 'output': str(Path(args.output).resolve())}, ensure_ascii=False))
        return 0 if result['complete'] else 2
    except (OSError, ValueError, RuntimeError, subprocess.CalledProcessError) as exc:
        print(f'执行失败：{exc}。如已生成目录，请保留并查看原因；没有完成标记的目录不算交付成功。', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())


