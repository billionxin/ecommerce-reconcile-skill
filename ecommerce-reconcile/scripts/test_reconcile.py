"""Generate synthetic Excel cases with Artifact Tool, then verify observable results."""
import copy
import json
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
from reconcile import SPECS, calculate, find_node, run, sha, money

ROOT = Path(__file__).resolve().parents[2]
VALIDATION = ROOT / '验证'
AS_OF = '2026-09-19'
BASE = {
 'orders': [['A001','2026-09-01','100.10','CNY'],['A002','2026-09-02','200.00','CNY']],
 'refunds': [['R001','A001','2026-09-03','20.00','成功','CNY']],
 'fees': [['F001','A001','2026-09-01','3.10','平台佣金','CNY'],['F002','A002','2026-09-02','6.00','平台佣金','CNY']],
 'costs': [['A001','50.00','10.00','0.00','CNY'],['A002','100.00','10.00','0.00','CNY']],
}


def cases():
    all_cases=[]
    def add(name, edit=None, code=None, profit=None, eligible=None, complete=None, **extra):
        data=copy.deepcopy(BASE)
        if edit: edit(data)
        all_cases.append(dict(name=name,data=data,code=code,profit=profit,eligible=eligible,complete=complete,**extra))
    def cell(k,r,c,v): return lambda d:d[k][r].__setitem__(c,v)
    add('01_正常四表',profit=10100,eligible=2,complete=True)
    add('02_重复订单相同内容',lambda d:d['orders'].append(d['orders'][0][:]),'DUPLICATE_ID',8400,1,False)
    add('03_重复订单冲突金额',lambda d:d['orders'].append(['A001','2026-09-01','999.00','CNY']),'DUPLICATE_ID',8400,1,False)
    add('04_缺少实收字段',code='SCHEMA',profit=0,eligible=0,complete=False,remove_header=('orders',2))
    add('05_金额空白',cell('orders',0,2,None),'INVALID_CELL',8400,1,False)
    add('06_数值型订单号',cell('orders',0,0,123),'INVALID_CELL',8400,1,False)
    add('07_前后空格不修剪',cell('orders',0,0,' A001 '),'INVALID_CELL',8400,1,False)
    add('08_重复退款单号',lambda d:d['refunds'].append(d['refunds'][0][:]),'DUPLICATE_ID',8400,1,False)
    add('09_多笔部分退款累加',lambda d:d['refunds'].append(['R002','A001','2026-09-04','10.01','成功','CNY']),profit=9099,eligible=2,complete=True)
    add('10_处理中退款不扣减',cell('refunds',0,4,'处理中'),'REFUND_PENDING',12100,2,True)
    add('11_失败退款不扣减',cell('refunds',0,4,'失败'),'REFUND_FAILED',12100,2,True)
    add('12_退款超过实收',cell('refunds',0,3,'100.11'),'REFUND_EXCEEDS_PAID',8400,1,False)
    add('13_孤立退款',cell('refunds',0,1,'X999'),'ORPHAN_ORDER',12100,2,False)
    add('14_重复费用单号',lambda d:d['fees'].append(d['fees'][0][:]),'DUPLICATE_ID',8400,1,False)
    add('15_多笔平台费用',lambda d:d['fees'].append(['F003','A001','2026-09-02','1.01','支付手续费','CNY']),profit=9999,eligible=2,complete=True)
    add('16_缺少费用不填零',lambda d:d['fees'].pop(0),'MISSING_FEE',8400,1,False)
    add('17_缺少成本不填零',lambda d:d['costs'].pop(0),'MISSING_COST',8400,1,False)
    add('18_重复成本',lambda d:d['costs'].append(d['costs'][0][:]),'DUPLICATE_ID',8400,1,False)
    add('19_负数费用',cell('fees',0,3,'-1.00'),'INVALID_CELL',8400,1,False)
    add('20_超过两位小数',cell('orders',0,2,'100.101'),'INVALID_CELL',8400,1,False)
    add('21_成本冲回超上限',cell('costs',0,3,'50.01'),'RECOVERY_EXCEEDS_COST',8400,1,False)
    add('22_负利润如实保留',cell('costs',0,1,'80.00'),'NEGATIVE_PROFIT',7100,2,True)
    add('23_费用超过实收提醒',cell('fees',0,3,'120.00'),'FEE_EXCEEDS_PAID',-1590,2,True)
    add('24_大额阈值提醒',cell('orders',0,2,'100000.00'),'LARGE_AMOUNT',10000090,2,True)
    add('25_公式金额拒绝',code='INVALID_CELL',profit=8400,eligible=1,complete=False,formula=('orders','C2','=100+0.1'))
    add('26_退款早于付款',cell('refunds',0,2,'2026-08-31'),'DATE_BEFORE_PAID',8400,1,False)
    add('27_无效日期',cell('orders',0,1,'2026-02-30'),'INVALID_CELL',8400,1,False)
    add('28_未来日期',cell('orders',0,1,'2026-09-20'),'INVALID_CELL',8400,1,False)
    add('29_混合币种拒绝',cell('fees',0,5,'USD'),'INVALID_CELL',8400,1,False)
    add('30_没有退款的空表',lambda d:d['refunds'].clear(),profit=12100,eligible=2,complete=True)
    add('31_空订单表',lambda d:d['orders'].clear(),'EMPTY_ORDERS',0,0,False)
    add('32_缺失文件',code='MISSING_FILE',profit=0,eligible=0,complete=False,missing='refunds')
    add('33_损坏工作簿',code='FILE_INVALID',profit=0,eligible=0,complete=False,corrupt='costs')
    add('34_额外工作表',code='SCHEMA',profit=0,eligible=0,complete=False,extra_sheet='orders')
    add('35_列顺序变化按表头关联',profit=10100,eligible=2,complete=True,reorder='orders')
    add('36_0_1加0_2精确分',lambda d:(d['fees'].__setitem__(0,['F001','A001','2026-09-01','0.10','平台佣金','CNY']),d['fees'].append(['F003','A001','2026-09-01','0.20','支付手续费','CNY'])),profit=10380,eligible=2,complete=True)
    def long_id(d):
        for k,rows in d.items():
            index=1 if k in ('refunds','fees') else 0
            for row in rows:
                if row[index]=='A001':row[index]='000012345678901234567890'
    add('37_长订单号和前导零',long_id,profit=10100,eligible=2,complete=True)
    add('38_千分位金额拒绝',cell('orders',0,2,'1,000.00'),'INVALID_CELL',8400,1,False)
    add('39_未知退款状态',cell('refunds',0,4,'已完成'),'INVALID_CELL',8400,1,False)
    add('40_退款不自动冲回成本',cell('refunds',0,3,'100.10'),'NEGATIVE_PROFIT',2090,2,True)
    add('41_明确成本冲回',cell('costs',0,3,'50.00'),profit=15100,eligible=2,complete=True)
    add('42_明确零费用',lambda d:(d['fees'][0].__setitem__(3,'0.00'),d['fees'][0].__setitem__(4,'无费用')),profit=10410,eligible=2,complete=True)
    add('43_无费用但非零',cell('fees',0,4,'无费用'),'INVALID_CELL',8400,1,False)
    add('44_空白行记录不丢失',lambda d:d['refunds'].insert(0,[None]*6),profit=10100,eligible=2,complete=True)
    add('45_跨订单重复退款号',lambda d:d['refunds'].append(['R001','A002','2026-09-04','10.00','成功','CNY']),'DUPLICATE_ID',0,0,False)
    add('46_零实收不支持',cell('orders',0,2,'0.00'),'INVALID_CELL',8400,1,False)
    add('47_布尔金额拒绝',cell('orders',0,2,True),'INVALID_CELL',8400,1,False)
    add('48_科学计数文本拒绝',cell('orders',0,2,'1e2'),'INVALID_CELL',8400,1,False)
    add('49_含时间文本拒绝',cell('orders',0,1,'2026-09-01 12:00:00'),'INVALID_CELL',8400,1,False)
    add('50_重复表头拒绝',code='SCHEMA',profit=0,eligible=0,complete=False,duplicate_header='orders')
    return all_cases


def demo(corrected=False):
    data={k:[] for k in SPECS}
    paid=[299,199,499,89,159,129,259,99,79,139]
    goods=[120,80,240,35,65,50,100,40,30,60]
    shipping=[12,10,15,8,10,10,12,8,8,10]
    for i,(p,g,s) in enumerate(zip(paid,goods,shipping),1):
        oid=f'O{1000+i}'
        data['orders'].append([oid,f'2026-09-{i:02}',str(p),'CNY'])
        data['costs'].append([oid,str(g),str(s),'240' if i==3 else '0','CNY'])
        data['fees'].append([f'F{1000+i}',oid,f'2026-09-{i:02}',f'{p*3//100}.{p*3%100:02}','平台佣金','CNY'])
    data['refunds']=[['R1002','O1002','2026-09-12','50','成功','CNY'],
                     ['R1003','O1003','2026-09-13','499','成功','CNY'],
                     ['R1005','O1005','2026-09-14','30','处理中','CNY'],
                     ['R1009','O1009','2026-09-15','29' if corrected else '99','成功','CNY'],
                     ['R1099','O1002' if corrected else 'O9999','2026-09-15','20','成功','CNY']]
    if not corrected:
        data['orders'].append(data['orders'][6][:])
        data['costs'][7][1]=None
        data['fees'].pop()
    return data


def specs_for(data, folder, case=None):
    result=[]
    case=case or {}
    for kind,(filename,headers) in SPECS.items():
        if case.get('missing')==kind or case.get('corrupt')==kind: continue
        headers=list(headers); rows=copy.deepcopy(data[kind])
        if case.get('remove_header',(None,None))[0]==kind:
            idx=case['remove_header'][1]; headers.pop(idx)
            for r in rows:r.pop(idx)
        if case.get('reorder')==kind:
            headers.reverse(); rows=[r[::-1] for r in rows]
        if case.get('duplicate_header')==kind:headers[-1]=headers[0]
        item=dict(path=str(folder/filename),headers=headers,rows=rows)
        if case.get('formula',(None,))[0]==kind:
            _,cell,formula=case['formula'];item['formulas']=[dict(cell=cell,formula=formula)]
        if case.get('extra_sheet')==kind:item['extra_sheet']='其他'
        result.append(item)
    return result


def prepare():
    specs=[]
    for case in cases():
        folder=VALIDATION/'模拟测试'/case['name'];folder.mkdir(parents=True,exist_ok=True)
        specs.extend(specs_for(case['data'],folder,case))
        if case.get('corrupt'):(folder/SPECS[case['corrupt']][0]).write_bytes(b'not-an-xlsx')
    specs.extend(specs_for({k:[] for k in SPECS},ROOT/'空白模板'))
    specs.extend(specs_for(demo(False),ROOT/'演示'/'待核对原始数据'))
    specs.extend(specs_for(demo(True),ROOT/'演示'/'凭证确认后的副本'))
    (VALIDATION/'模拟生成清单.json').write_text(json.dumps(specs,ensure_ascii=False),encoding='utf-8')
    print(f'Prepared {len(specs)} workbooks')


def verify():
    rows=[]
    for c in cases():
        folder=VALIDATION/'模拟测试'/c['name']
        before={p.name:sha(p.read_bytes()) for p in folder.glob('*.xlsx')}
        try:
            r=calculate(folder,AS_OF)
            codes={i['code'] for i in r['issues']}
            if c['code']:assert c['code'] in codes,(c['code'],codes)
            if c['profit'] is not None:assert r['totals_cents']['profit']==c['profit'],(r['totals_cents']['profit'],c['profit'])
            if c['eligible'] is not None:assert r['eligible_orders']==c['eligible']
            if c['complete'] is not None:assert r['complete']==c['complete']
            assert before=={p.name:sha(p.read_bytes()) for p in folder.glob('*.xlsx')},'原文件改变'
            for kind,source in r['sources'].items():
                assert sum(r['row_counts'][kind].values())==source['data_rows'],'行去向不守恒'
            for d in r['details']:
                if d['status']=='待核实':assert d['amounts_cents'] is None
            if c['name'].startswith('36'):assert r['details'][0]['amounts_cents']['fee']==30
            if c['name'].startswith('37'):assert r['details'][0]['order_id']=='000012345678901234567890'
            if c['name'].startswith('44'):assert 2 in r['sources']['refunds']['blank_rows'] and r['ledger']['refunds'][0]['row']==3
            (folder/'预期与实际.json').write_text(json.dumps(dict(expected={k:v for k,v in c.items() if k!='data'},
                actual=dict(profit_cents=r['totals_cents']['profit'],eligible_orders=r['eligible_orders'],complete=r['complete'],codes=sorted(codes))),ensure_ascii=False,indent=2),encoding='utf-8')
            rows.append([c['name'],'通过',str(c['profit']),str(r['totals_cents']['profit'])])
        except Exception as exc:
            rows.append([c['name'],'失败',str(c['profit']),repr(exc)])
    # End-to-end checks use the actual produced customer workbooks and archived originals.
    checks=[]
    for label,expected,eligible in [('发现问题',36878,6),('确认后重跑',61050,10)]:
        folder=ROOT/'演示'/label
        r=json.loads((folder/'审计记录.json').read_text(encoding='utf-8'))
        assert r['totals_cents']['profit']==expected
        assert r['eligible_orders']==eligible
        w=openpyxl.load_workbook(folder/'经营利润报表.xlsx',data_only=False)
        assert w.sheetnames==['经营汇总','订单明细','异常清单','计算说明']
        assert money(w['经营汇总']['B20'].value)==expected
        for ws in w:
            for row in ws:
                for cell in row:assert cell.data_type not in ('f','e'),f'导出含公式或错误值 {ws.title}!{cell.coordinate}'
        for d in r['details']:
            matches=[row for row in w['订单明细'].iter_rows(min_row=2,values_only=True) if row[0]==d['order_id']]
            assert len(matches)==1
            if d['amounts_cents'] is None:assert all(v is None for v in matches[0][3:10])
        source_dir = ROOT/'演示'/('待核对原始数据' if label == '发现问题' else '凭证确认后的副本')
        for s in r['sources'].values():
            assert sha((folder/'原始数据'/s['file']).read_bytes())==s['sha256']
            assert sha((source_dir/s['file']).read_bytes())==s['sha256']
        for name,digest in json.loads((folder/'交付校验.json').read_text(encoding='utf-8')).items():
            assert sha((folder/name).read_bytes())==digest
        w.close()
        checks.append(label+'：Excel 汇总、明细空值、无公式、原文件哈希与导出哈希核对通过')
    repeated=calculate(ROOT/'演示'/'凭证确认后的副本',AS_OF)
    assert repeated['totals_cents']['profit']==61050
    try:
        run(ROOT/'演示'/'凭证确认后的副本',ROOT/'演示'/'确认后重跑',AS_OF)
    except FileExistsError:checks.append('已有输出目录覆盖被拒绝')
    else:raise AssertionError('覆盖保护失效')
    checks.append('同一输入重跑，金额结果一致')
    failed=sum(r[1]=='失败' for r in rows)
    report=f'# 模拟数据测试报告\n\n运行时间（UTC）：{datetime.now(timezone.utc).isoformat()}\n\n共 {len(rows)} 组 Excel 场景，{len(rows)-failed} 组通过，{failed} 组失败。金额列单位：分。\n\n|场景|结果|预期贡献|实际贡献或错误|\n|---|---|---:|---:|\n'
    report+='\n'.join('|'+ '|'.join(row)+'|' for row in rows)+'\n\n端到端核验：\n\n'+'\n'.join('- '+c for c in checks)
    report+='\n\n测试采用自行构造的模拟数据，不代表真实客户账目。所有测试输入及逐组预期/实际保存在 模拟测试/。\n'
    (VALIDATION/'测试报告.md').write_text(report,encoding='utf-8')
    print(f'{len(rows)-failed}/{len(rows)} cases passed; {len(checks)} integration checks passed')
    for row in rows:
        if row[1]=='失败':print(row)
    return failed


if __name__=='__main__':
    if '--prepare' in sys.argv:prepare()
    else:sys.exit(verify())


